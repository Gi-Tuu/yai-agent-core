"""把任意 markdown 构建成带学术封面、自动目录与页码的 Word，并可选转 PDF（比赛材料通用）。

用法：
  .venv/Scripts/python.exe scripts/build_competition_docx.py \\
      docs/competitions/shanghai/作品介绍.md docs/exports/作品介绍.docx --pdf

约定：
- markdown 第一个一级标题作为封面主标题；
- 第一个 h1 之后、第一个 h2 之前的块作为封面内容（引用块作副标题，表格作元数据表）；
- 封面后插入可刷新的 Word 目录（TOC，1-2 级）；正文 h2 为章、另起一页；
- 图片以其 alt 文本作为居中图注；
- --pdf 时调用本机 LibreOffice（soffice）无头转换，并设置 updateFields 以更新目录页码。
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
from pathlib import Path

# 允许直接运行（python scripts/build_competition_docx.py）时 import 同目录脚本
sys.path.insert(0, str(Path(__file__).resolve().parent))

from build_walkthrough_docx import (  # noqa: E402
    add_inline,
    parse_md,
    parse_table,
    preserve_spaces,
    set_east_asian,
    shade_paragraph,
)
from docx import Document  # noqa: E402
from docx.enum.style import WD_STYLE_TYPE  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

SOFFICE_CANDIDATES = (
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    "/usr/bin/soffice",
    "/opt/homebrew/bin/soffice",
)


def find_soffice() -> str | None:
    found = shutil.which("soffice")
    if found:
        return found
    for path in SOFFICE_CANDIDATES:
        if Path(path).exists():
            return path
    return None


def configure_styles(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    set_east_asian(normal, "Arial", "宋体", Pt(12))
    normal.paragraph_format.line_spacing = 1.5
    ppr = normal.element.get_or_add_pPr()
    ind = OxmlElement("w:ind")
    ind.set(qn("w:firstLineChars"), "200")
    ppr.append(ind)

    for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        st = doc.styles[name]
        set_east_asian(st, "Arial", "黑体", Pt(size), bold=True, color="000000")
    title_style = doc.styles["Title"]
    set_east_asian(title_style, "Arial", "黑体", Pt(24), bold=True)
    title_bdr = title_style.element.get_or_add_pPr().find(qn("w:pBdr"))
    if title_bdr is not None:
        title_style.element.get_or_add_pPr().remove(title_bdr)

    # 目录页中“目录”字样（不被 TOC 收录，故不使用 Heading 样式）
    toc_head = doc.styles.add_style("TOCHead", WD_STYLE_TYPE.PARAGRAPH)
    toc_head.base_style = doc.styles["Normal"]
    set_east_asian(toc_head, "Arial", "黑体", Pt(18), bold=True, color="000000")
    thpr = toc_head.element.get_or_add_pPr()
    thind = OxmlElement("w:ind")
    thind.set(qn("w:firstLine"), "0")
    thind.set(qn("w:firstLineChars"), "0")
    thpr.append(thind)

    code_style = doc.styles.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
    code_style.base_style = doc.styles["Normal"]
    set_east_asian(code_style, "Consolas", "宋体", Pt(10))
    code_style.paragraph_format.line_spacing = 1.15
    code_style.paragraph_format.left_indent = Cm(0.4)
    code_style.paragraph_format.space_before = Pt(0)
    code_style.paragraph_format.space_after = Pt(0)
    cpr = code_style.element.get_or_add_pPr()
    cind = OxmlElement("w:ind")
    cind.set(qn("w:firstLine"), "0")
    cind.set(qn("w:firstLineChars"), "0")
    cpr.append(cind)


def add_footer_pagenum(doc: Document) -> None:
    footer_par = doc.sections[0].footer.paragraphs[0]
    footer_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fb = OxmlElement("w:fldChar")
    fb.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = " PAGE "
    fe = OxmlElement("w:fldChar")
    fe.set(qn("w:fldCharType"), "end")
    run = footer_par.add_run()._r
    for el in (fb, it, fe):
        run.append(el)


def enable_update_fields(doc: Document) -> None:
    """让 Word/LibreOffice 在加载时更新所有域（用于生成目录页码）。"""
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    doc.settings.element.append(upd)


def add_cover_meta_table(doc: Document, payload) -> None:
    """封面元数据表：两列，首列浅底加粗，无表头行。"""
    _, body = parse_table(payload)
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row in body:
        cells = table.add_row().cells
        for j, val in enumerate(row[:2]):
            cells[j].text = ""
            par = cells[j].paragraphs[0]
            par.paragraph_format.first_line_indent = Cm(0)
            run = par.add_run(val)
            run.font.size = Pt(10.5)
            if j == 0:
                run.bold = True
                tcpr = cells[j]._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), "EDF1F6")
                tcpr.append(shd)
    # 列宽
    for row in table.rows:
        row.cells[0].width = Cm(4.2)
        row.cells[1].width = Cm(11.3)


def add_toc(doc: Document) -> None:
    head = doc.add_paragraph(style="TOCHead")
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    head.add_run("目　录")
    par = doc.add_paragraph()
    par.paragraph_format.first_line_indent = Cm(0)
    run = par.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-2" \h \z \u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "（目录将在打开文档或转换时自动生成）"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, separate, placeholder, end):
        run._r.append(el)


def add_figure(doc: Document, payload: dict, base_dir: Path) -> None:
    """插入居中图片，并以 alt 作为正式图注（10.5pt 深灰居中）。"""
    img_path = Path(payload["path"])
    if not img_path.is_absolute():
        img_path = base_dir / img_path
    pic_par = doc.add_paragraph()
    pic_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_par.paragraph_format.first_line_indent = Cm(0)
    pic_par.paragraph_format.keep_with_next = True
    pic_par.paragraph_format.keep_together = True
    if img_path.exists():
        pic_par.add_run().add_picture(str(img_path), width=Cm(15.2))
    else:
        r = pic_par.add_run(f"[图片缺失：{img_path}]")
        r.font.color.rgb = RGBColor.from_string("C0392B")
    alt = payload.get("alt")
    if alt:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.first_line_indent = Cm(0)
        cap.paragraph_format.keep_together = True
        cr = cap.add_run(alt)
        cr.font.size = Pt(10.5)
        cr.font.color.rgb = RGBColor.from_string("333333")
        cr.bold = True


def build_docx(md_path: Path, out_docx: Path) -> None:
    blocks = parse_md(md_path.read_text(encoding="utf-8"))

    title = md_path.stem
    cover_blocks: list[tuple[str, object]] = []
    body_blocks: list[tuple[str, object]] = []
    h1_seen = False
    in_cover = True
    for kind, payload in blocks:
        if kind == "h1" and not h1_seen:
            title = str(payload)
            h1_seen = True
            continue
        if in_cover:
            if kind == "h2":
                in_cover = False
                body_blocks.append((kind, payload))
            else:
                cover_blocks.append((kind, payload))
        else:
            body_blocks.append((kind, payload))

    doc = Document()
    configure_styles(doc)
    add_footer_pagenum(doc)
    enable_update_fields(doc)

    # ---------------- 封面
    for _ in range(5):
        doc.add_paragraph()
    t = doc.add_paragraph(style="Title")
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.first_line_indent = Cm(0)
    t.add_run(title)
    quote_n = 0
    for kind, payload in cover_blocks:
        if kind == "quote":
            sub = doc.add_paragraph()
            sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            sub.paragraph_format.first_line_indent = Cm(0)
            r = sub.add_run(str(payload))
            if quote_n == 0:
                r.font.size = Pt(15)
                r.bold = True
                r.font.color.rgb = RGBColor.from_string("1F3A5F")
            else:
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor.from_string("4A4A4A")
            quote_n += 1
        elif kind == "table":
            for _ in range(3):
                doc.add_paragraph()
            add_cover_meta_table(doc, payload)
        elif kind == "p":
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(par, str(payload))
    doc.add_page_break()

    # ---------------- 目录
    add_toc(doc)
    doc.add_page_break()

    # 有序列表独立编号（每个列表从 1 开始）
    numbering_el = doc.part.numbering_part.element
    ln_numpr = doc.styles["List Number"].element.find(qn("w:pPr")).find(qn("w:numPr"))
    base_num_id = int(ln_numpr.find(qn("w:numId")).get(qn("w:val")))
    base_num = next(
        n for n in numbering_el.findall(qn("w:num"))
        if int(n.get(qn("w:numId"))) == base_num_id
    )
    abstract_num_id = int(base_num.find(qn("w:abstractNumId")).get(qn("w:val")))

    def new_restart_num_id() -> int:
        existing = [int(n.get(qn("w:numId"))) for n in numbering_el.findall(qn("w:num"))]
        num_id = max(existing) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        aref = OxmlElement("w:abstractNumId")
        aref.set(qn("w:val"), str(abstract_num_id))
        num.append(aref)
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:startOverride")
        start.set(qn("w:val"), "1")
        override.append(start)
        num.append(override)
        numbering_el.append(num)
        return num_id

    first_h2 = True
    for kind, payload in body_blocks:
        if kind == "h2":
            # 章标题另起一页（第一个 h2 紧随目录页，不再额外加页）
            if not first_h2:
                doc.add_page_break()
            first_h2 = False
            doc.add_paragraph(str(payload), style="Heading 2")
        elif kind.startswith("h"):
            doc.add_paragraph(str(payload), style=f"Heading {kind[1]}")
        elif kind == "p":
            par = doc.add_paragraph()
            add_inline(par, str(payload))
        elif kind == "quote":
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.6)
            add_inline(par, str(payload))
        elif kind == "ul":
            for item in payload:  # type: ignore[assignment]
                par = doc.add_paragraph(style="List Bullet")
                add_inline(par, item)
        elif kind == "ol":
            num_id = new_restart_num_id()
            for item in payload:  # type: ignore[assignment]
                par = doc.add_paragraph(style="List Number")
                ppr = par._p.get_or_add_pPr()
                numpr = OxmlElement("w:numPr")
                ilvl = OxmlElement("w:ilvl")
                ilvl.set(qn("w:val"), "0")
                nid = OxmlElement("w:numId")
                nid.set(qn("w:val"), str(num_id))
                numpr.append(ilvl)
                numpr.append(nid)
                ppr.append(numpr)
                add_inline(par, item)
        elif kind == "code":
            for code_line in str(payload).split("\n"):
                par = doc.add_paragraph(style="CodeBlock")
                run = par.add_run(code_line if code_line else " ")
                preserve_spaces(run)
                shade_paragraph(par, "F4F5F7")
        elif kind == "table":
            header, body = parse_table(payload)  # type: ignore[arg-type]
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = table.rows[0].cells
            for j, h in enumerate(header):
                hdr[j].text = ""
                run = hdr[j].paragraphs[0].add_run(h)
                run.bold = True
                run.font.size = Pt(10.5)
                tcpr = hdr[j]._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), "D9D9D9")
                tcpr.append(shd)
            for row in body:
                cells = table.add_row().cells
                for j, val in enumerate(row):
                    if j >= len(cells):
                        break
                    cells[j].text = ""
                    add_inline(cells[j].paragraphs[0], val, base_size=Pt(10.5))
            doc.add_paragraph()
        elif kind == "image":
            add_figure(doc, payload, md_path.parent)

    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_docx)
    print(f"saved: {out_docx}")


def convert_pdf(out_docx: Path) -> Path | None:
    soffice = find_soffice()
    if not soffice:
        print("未找到 LibreOffice（soffice），跳过 PDF；请安装后手动转换。")
        return None
    subprocess.run(
        [
            soffice, "--headless", "--convert-to", "pdf",
            "--outdir", str(out_docx.parent), str(out_docx),
        ],
        check=True,
        timeout=240,
    )
    pdf = out_docx.with_suffix(".pdf")
    print(f"saved: {pdf}")
    return pdf


def main() -> None:
    parser = argparse.ArgumentParser(description="markdown → 比赛材料 Word/PDF")
    parser.add_argument("md", type=Path)
    parser.add_argument("docx", type=Path)
    parser.add_argument("--pdf", action="store_true", help="同时用 LibreOffice 转 PDF")
    args = parser.parse_args()

    build_docx(args.md.resolve(), args.docx.resolve())
    if args.pdf:
        convert_pdf(args.docx.resolve())


if __name__ == "__main__":
    main()
